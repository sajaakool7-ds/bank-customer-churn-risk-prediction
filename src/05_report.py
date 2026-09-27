import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage

risk_df = pd.read_csv('/home/claude/churn_project/outputs/customer_risk_report.csv')

# ---------- Real metrics from the executed notebook ----------
roc_auc_best = 0.8699
pr_auc_best = 0.7135
overall_test_churn = risk_df['ActualExited'].mean()

comparison = pd.DataFrame({
    'Model': ['XGBoost (tuned)', 'XGBoost + SMOTE', 'Random Forest (tuned)', 'Logistic Regression'],
    'ROC-AUC': [0.8699, 0.8648, 0.8612, 0.8502],
    'PR-AUC':  [0.7135, 0.7153, 0.7063, 0.6887],
})

tier_validation = risk_df.groupby('RiskTier')['ActualExited'].agg(['mean', 'count']) \
    .reindex(['Critical', 'High', 'Medium', 'Low'])

shap_importance = [
    ("NumOfProducts", 0.729),
    ("Age", 0.637),
    ("IsActiveMember", 0.285),
    ("Gender (Male)", 0.233),
    ("IsGermany", 0.200),
]

wb = Workbook()

# ---------- Sheet 1: Executive Summary ----------
ws = wb.active
ws.title = "Executive Summary"
header_font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
sub_font = Font(name="Arial", bold=True, size=11)
normal_font = Font(name="Arial", size=10)
fill_header = PatternFill("solid", fgColor="1F4E78")
fill_critical = PatternFill("solid", fgColor="C00000")
fill_high = PatternFill("solid", fgColor="ED7D31")
fill_medium = PatternFill("solid", fgColor="FFC000")
fill_low = PatternFill("solid", fgColor="70AD47")

ws['A1'] = "Bank Customer Churn — Risk & Retention Decision System"
ws['A1'].font = header_font
ws.merge_cells('A1:D1')
ws['A1'].fill = fill_header
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 28

ws['A3'] = "Model: XGBoost (tuned via RandomizedSearchCV, scale_pos_weight for imbalance)"
ws['A3'].font = sub_font
ws['A4'] = f"Test set ROC-AUC: {roc_auc_best}  |  PR-AUC: {pr_auc_best}  |  Test size: {len(risk_df):,} customers (20% holdout)"
ws['A4'].font = normal_font
ws['A5'] = f"Overall test churn rate: {overall_test_churn:.1%}"
ws['A5'].font = normal_font

ws['A7'] = "Model Comparison"
ws['A7'].font = sub_font
ws.append([])
ws.append(['Model', 'ROC-AUC', 'PR-AUC'])
for c in ws[9]:
    c.font = Font(name="Arial", bold=True)
    c.fill = PatternFill("solid", fgColor="D9E1F2")
for _, row in comparison.iterrows():
    ws.append([row['Model'], row['ROC-AUC'], row['PR-AUC']])

r0 = ws.max_row + 2
ws[f'A{r0}'] = "Risk Tier Distribution (on holdout test set)"
ws[f'A{r0}'].font = sub_font

headers = ['Risk Tier', '# Customers', 'Actual Churn Rate (validation)']
ws.append(headers)
hdr_row = ws.max_row
for c in ws[hdr_row]:
    c.font = Font(name="Arial", bold=True)
    c.fill = PatternFill("solid", fgColor="D9E1F2")

fills = {'Critical': fill_critical, 'High': fill_high, 'Medium': fill_medium, 'Low': fill_low}
for tier, row in tier_validation.iterrows():
    ws.append([tier, int(row['count']), round(row['mean'], 3)])
    ws.cell(row=ws.max_row, column=1).fill = fills[tier]
    ws.cell(row=ws.max_row, column=1).font = Font(name="Arial", bold=True, color="FFFFFF")

r1 = ws.max_row + 2
ws[f'A{r1}'] = "Top Churn Drivers (global SHAP importance)"
ws[f'A{r1}'].font = sub_font
drivers = [
    ("1. Number of products", "3-4 products correlates with very high churn — signals a service/cross-sell conflict, not loyalty"),
    ("2. Age", "Churn rises sharply for older customers, especially 45-60"),
    ("3. Active membership status", "Inactive members churn noticeably more than active ones"),
    ("4. Gender", "Female customers churn more than male customers"),
    ("5. Geography — Germany", "Germany churns at a notably higher rate than France/Spain"),
]
r = r1 + 1
for title, desc in drivers:
    ws[f'A{r}'] = title
    ws[f'A{r}'].font = Font(name="Arial", bold=True, size=10)
    ws[f'B{r}'] = desc
    ws[f'B{r}'].font = normal_font
    ws.merge_cells(f'B{r}:D{r}')
    r += 1

try:
    img = XLImage('/home/claude/churn_project/outputs/shap_summary.png')
    img.width = 560
    img.height = 460
    ws.add_image(img, f'A{r + 2}')
except Exception as e:
    print("image embed skipped:", e)

for col, w in zip(['A', 'B', 'C', 'D'], [30, 55, 22, 22]):
    ws.column_dimensions[col].width = w

# ---------- Sheet 2: Action Plan ----------
ws2 = wb.create_sheet("Action Plan")
ws2['A1'] = "Recommended Action Plan by Risk Tier & Root Cause"
ws2['A1'].font = header_font
ws2.merge_cells('A1:C1')
ws2['A1'].fill = fill_header
ws2.row_dimensions[1].height = 24

playbook = [
    ("NumOfProducts ≥ 3 (any tier)", "Immediate review by Relationship Manager",
     "High product count here signals conflicting fees/overlapping services, not upsell success — escalate to service recovery, not sales"),
    ("Top reason: IsActiveMember / InactiveSenior", "Re-engagement campaign",
     "Inactive customers disengage quietly; a proactive nudge outperforms passive marketing"),
    ("Top reason: Age", "Age-specific offer and proactive call",
     "Churn concentrates in older age brackets — a tailored offer plus a human touchpoint works better than generic messaging"),
    ("Top reason: Germany / IsGermany", "Review pricing and service in Germany",
     "Germany's churn is structurally higher — investigate pricing/competitive pressure in that market"),
    ("No specific top driver", "Regular follow-up and general retention offer",
     "Default action when no single dominant SHAP driver stands out — keep the customer engaged with routine outreach"),
]
ws2.append(["Condition (from TopReason1 / NumOfProducts)", "Recommended Action", "Rationale"])
for c in ws2[2]:
    c.font = Font(name="Arial", bold=True)
    c.fill = PatternFill("solid", fgColor="D9E1F2")
for row in playbook:
    ws2.append(row)

r2 = ws2.max_row + 2
ws2[f'A{r2}'] = "Distribution of customers per action (test set)"
ws2[f'A{r2}'].font = sub_font
ws2.append(["Recommended Action", "Customer Count", ""])
hdr2 = ws2.max_row
for c in ws2[hdr2]:
    c.font = Font(name="Arial", bold=True)
    c.fill = PatternFill("solid", fgColor="D9E1F2")

action_counts = risk_df['RecommendedAction'].value_counts()
for action, count in action_counts.items():
    ws2.append([action, int(count), ""])

for col, w in zip(['A', 'B', 'C'], [40, 40, 60]):
    ws2.column_dimensions[col].width = w
for row in ws2.iter_rows(min_row=3, max_row=len(playbook) + 2):
    for cell in row:
        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(wrap_text=True, vertical='top')

# ---------- Sheet 3: Customer Risk Report (full data, 2000 rows) ----------
ws3 = wb.create_sheet("Customer Risk Report")
cols = ['CustomerId', 'ChurnProbability', 'RiskTier', 'ActualExited',
        'TopReason1', 'TopReason2', 'TopReason3', 'RecommendedAction']
ws3.append(cols)
for c in ws3[1]:
    c.font = Font(name="Arial", bold=True, color="FFFFFF")
    c.fill = fill_header
for _, row in risk_df[cols].iterrows():
    ws3.append(list(row))
widths = [12, 15, 10, 12, 25, 25, 25, 45]
for col_letter, w in zip('ABCDEFGH', widths):
    ws3.column_dimensions[col_letter].width = w
for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
    for cell in row:
        cell.font = Font(name="Arial", size=9)

wb.save('/home/claude/churn_project/outputs/churn_retention_report.xlsx')
print("Saved churn_retention_report.xlsx")
